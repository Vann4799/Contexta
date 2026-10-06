from io import BytesIO

import fitz
from docx import Document

from worker.extraction import DocumentTextExtractor


def test_extracts_pdf_text_by_page() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Hello PDF page")
    content = pdf.tobytes()
    pdf.close()

    extracted = DocumentTextExtractor().extract(content, "pdf")

    assert extracted["pages"] == [{"page_number": 1, "text": "Hello PDF page"}]


def test_extracts_docx_paragraph_text() -> None:
    document = Document()
    document.add_paragraph("Hello DOCX")
    document.add_paragraph("Second paragraph")
    buffer = BytesIO()
    document.save(buffer)

    extracted = DocumentTextExtractor().extract(buffer.getvalue(), "docx")

    assert extracted["pages"] == [
        {"page_number": 1, "text": "Hello DOCX\nSecond paragraph"}
    ]


def test_pdf_lines_larger_than_body_become_headings() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "BAB IV HASIL", fontsize=20)
    page.insert_text((72, 100), "isi satu", fontsize=10)
    page.insert_text((72, 120), "isi dua", fontsize=10)
    content = pdf.tobytes()
    pdf.close()

    extracted = DocumentTextExtractor().extract(content, "pdf")

    lines = extracted["pages"][0]["text"].splitlines()
    assert lines[0] == "# BAB IV HASIL"
    assert "isi satu" in lines
    assert not any(line.startswith("#") for line in lines[1:])


def test_docx_heading_styles_become_heading_markers() -> None:
    document = Document()
    document.add_heading("BAB I", level=1)
    document.add_heading("1.1 Latar Belakang", level=2)
    document.add_paragraph("Paragraf biasa.")
    buffer = BytesIO()
    document.save(buffer)

    extracted = DocumentTextExtractor().extract(buffer.getvalue(), "docx")

    assert extracted["pages"][0]["text"] == (
        "# BAB I\n## 1.1 Latar Belakang\nParagraf biasa."
    )
